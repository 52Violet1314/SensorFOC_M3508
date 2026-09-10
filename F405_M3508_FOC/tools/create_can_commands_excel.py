from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


OUT = Path(__file__).resolve().parents[1] / "CAN_Control_Commands.xlsx"


def signed_bytes(raw: int):
    raw &= 0xFFFF
    return [(raw >> 8) & 0xFF, raw & 0xFF]


def add_sheet(wb, title, mode, values, scale, unit, decode):
    ws = wb.create_sheet(title)
    ws.append(["组号", "目标值", "单位", "协议原始值", "CAN ID", "DLC", "数据字节0", "数据字节1", "数据字节2", "数据字节3", "数据字节4", "数据字节5", "数据字节6", "数据字节7", "完整指令（单元格复制）"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for index, value in enumerate(values, 1):
        raw = int(round(value * scale))
        b1, b2 = signed_bytes(raw)
        frame = [mode, b1, b2, 0, 0, 0, 0, 0]
        shown_value = decode(raw)
        byte_text = [f"{byte:02X}" for byte in frame]
        full_command = "0x91 8 " + " ".join(byte_text)
        ws.append([index, shown_value, unit, raw, "0x91", 8] + byte_text + [full_command])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["O"].width = 42
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "0.000"
        ws.cell(row, 4).number_format = "0"


wb = Workbook()
intro = wb.active
intro.title = "协议说明"
intro.append(["项目", "说明"])
intro.append(["接收 CAN ID", "0x91"])
intro.append(["数据格式", "DLC=8，数据字节1~2为有符号大端 int16"])
intro.append(["电流模式", "数据字节0=01；实际电流 = 原始值 / 1000 A"])
intro.append(["速度模式", "数据字节0=02；实际速度 = 原始值 rpm"])
intro.append(["位置模式", "数据字节0=03；实际位置 = 原始值 × pi / 30000 rad，单圈"])
intro.append(["保留字节", "数据字节3~7暂时填 00"])
for cell in intro[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
for col, width in ((1, 18), (2, 80)):
    intro.column_dimensions[get_column_letter(col)].width = width

current_values = [-10.0 + 20.0 * i / 9.0 for i in range(10)]
speed_values = [-5000.0 + 10000.0 * i / 9.0 for i in range(10)]
position_values = [-3.14 + 6.28 * i / 9.0 for i in range(10)]

add_sheet(wb, "电流指令", 0x01, current_values, 1000.0, "A", lambda raw: raw / 1000.0)
add_sheet(wb, "速度指令", 0x02, speed_values, 1.0, "rpm", lambda raw: float(raw))
add_sheet(wb, "位置指令", 0x03, position_values, 30000.0 / 3.141592653589793, "rad", lambda raw: raw * 3.141592653589793 / 30000.0)

wb.save(OUT)
print(OUT)
